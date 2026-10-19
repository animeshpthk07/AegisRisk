"""Strict values-only XLSX decoding with bounded archive and XML parsing."""

import csv
import io
import re
from decimal import Decimal
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXPANDED = 4 * 1024 * 1024
MAX_MEMBER = 2 * 1024 * 1024
CELL_REF = re.compile(r"[A-J]([1-9][0-9]{0,3})", re.ASCII)


def check_archive(raw: bytes, max_rows: int):
    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if not entries or len(entries) > 64:
            raise HTTPException(413, "XLSX archive must contain at most 64 entries")
        names = set()
        expanded = 0
        for entry in entries:
            name = entry.filename
            lower = name.lower()
            if (lower in names or name.startswith("/") or ".." in name.split("/")
                    or "\\" in name or entry.flag_bits & 1
                    or entry.compress_type not in (ZIP_STORED, ZIP_DEFLATED)):
                raise HTTPException(422, "Unsupported or ambiguous XLSX archive")
            names.add(lower)
            if (not lower.endswith((".xml", ".rels"))
                    or any(part in lower for part in ("externallinks/", "embeddings/", "vba"))):
                raise HTTPException(422, "Use a plain XLSX without macros or embedded content")
            expanded += entry.file_size
            if (entry.file_size > MAX_MEMBER or expanded > MAX_EXPANDED
                    or entry.file_size > max(entry.compress_size, 1) * 100):
                raise HTTPException(413, "XLSX expanded content exceeds safety limits")
            with archive.open(entry) as member:
                payload = member.read(MAX_MEMBER + 1)
            if len(payload) != entry.file_size or len(payload) > MAX_MEMBER:
                raise HTTPException(413, "XLSX member exceeds safety limits")
            root = fromstring(payload, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            coordinates = set()
            for node in root.iter():
                tag = node.tag.rsplit("}", 1)[-1]
                if (tag in {"f", "definedName", "hyperlink", "mergeCell", "autoFilter"}
                        or node.get("TargetMode", "").lower() == "external"
                        or any("macroenabled" in value.lower() for value in node.attrib.values())):
                    raise HTTPException(
                        422, "Use values only; formulas, links and filters are rejected",
                    )
                if tag in {"row", "col", "sheet"} and (
                    node.get("hidden") in {"1", "true"}
                    or node.get("state", "visible") != "visible"
                ):
                    raise HTTPException(422, "Hidden rows, columns and sheets are not supported")
                if tag == "c" and lower.startswith("xl/worksheets/"):
                    cell_ref = node.get("r", "")
                    match = CELL_REF.fullmatch(cell_ref)
                    if not match or int(match[1]) > max_rows + 1:
                        raise HTTPException(413, "XLSX must fit A1:J1001 including the header")
                    if cell_ref in coordinates:
                        raise HTTPException(422, "Duplicate XLSX cell coordinates")
                    coordinates.add(cell_ref)


def xlsx_to_csv(raw: bytes, max_bytes: int, max_rows: int):
    if not raw or len(raw) > max_bytes:
        raise HTTPException(413, "Choose a non-empty XLSX no larger than 256 KiB")
    workbook = None
    try:
        check_archive(raw, max_rows)
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
            raise HTTPException(422, "XLSX must contain exactly one visible worksheet")
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        records = []
        for row in sheet.iter_rows(min_row=1, max_row=max_rows + 1, min_col=1, max_col=10):
            values = []
            for cell in row:
                value = cell.value
                if cell.data_type in {"f", "e", "b", "d"}:
                    raise HTTPException(422, "Formula, error, boolean and date cells are rejected")
                if value is None:
                    values.append("")
                elif isinstance(value, str):
                    values.append(value)
                elif isinstance(value, (int, float)):
                    number = Decimal(str(value))
                    if not number.is_finite():
                        raise HTTPException(422, "XLSX numeric cells must be finite")
                    values.append(format(number, "f"))
                else:
                    raise HTTPException(422, "Use text and numeric values only in XLSX cells")
            records.append(values if any(values) else [])
        while records and not records[-1]:
            records.pop()
        output = io.StringIO(newline="")
        csv.writer(output, lineterminator="\n").writerows(records)
        normalized = output.getvalue().encode("utf-8")
        if len(normalized) > max_bytes:
            raise HTTPException(413, "Normalized XLSX values exceed the 256 KiB limit")
        return normalized
    except HTTPException:
        raise
    except (BadZipFile, DefusedXmlException, ParseError, InvalidFileException,
            ValueError, KeyError, TypeError, IndexError, OSError, EOFError):
        raise HTTPException(422, "Malformed or unsupported XLSX workbook") from None
    finally:
        if workbook is not None:
            workbook.close()
