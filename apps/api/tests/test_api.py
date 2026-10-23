import csv
import io
import json
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, delete, func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import aegisrisk.main as api_module
from aegisrisk.database import Base, get_session
from aegisrisk.main import app
from aegisrisk.models import Facility, SavedPortfolio


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(
            Facility(
                id="A",
                borrower="Fictional A",
                industry="Energy",
                geography="West",
                rating="B",
                pd=Decimal(".1"),
                lgd=Decimal(".5"),
                ead=Decimal(1000),
            )
        )
        session.commit()

    def session_override():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_session] = session_override
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    engine.dispose()


def test_health_and_ready(client):
    assert client.get("/api/v1/health").json()["external_ai"] is False
    assert client.get("/api/v1/ready").json()["database"] == "connected"


def test_portfolio(client):
    response = client.get("/api/v1/portfolio")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.json()["summary"]["total_expected_loss"] == 50
    assert response.json()["synthetic"] is True


def test_pagination_search(client):
    assert client.get("/api/v1/facilities?q=Fictional&limit=1").json()["total"] == 1
    assert client.get("/api/v1/facilities?q=%25").json()["total"] == 0
    assert client.get("/api/v1/facilities?offset=1").json()["items"] == []
    assert client.get("/api/v1/facilities?limit=1000").status_code == 422
    assert client.get("/api/v1/facilities?offset=-1").status_code == 422


def test_stress(client):
    result = client.post(
        "/api/v1/stress", json={"pd_multiplier": 2, "lgd_multiplier": 1, "ead_multiplier": 1}
    ).json()
    assert result["baseline_el"] == 50
    assert result["stressed_el"] == 100
    assert result["delta_pct"] == 1


def test_stress_probability_cap(client):
    result = client.post(
        "/api/v1/stress", json={"pd_multiplier": 10, "lgd_multiplier": 10, "ead_multiplier": 1}
    ).json()
    assert result["stressed_el"] == 1000


@pytest.mark.parametrize(
    "payload",
    [
        {"pd_multiplier": -1},
        {"pd_multiplier": "NaN"},
        {"pd_multiplier": 11},
        {"unknown": 1},
        {"industry": "Not real"},
    ],
)
@pytest.mark.parametrize("resource", ["stress", "scenario-report", "scenario-report-pdf"])
def test_invalid_stress(client, payload, resource):
    assert client.post(f"/api/v1/{resource}", json=payload).status_code == 422


def test_grounded_explanation(client):
    result = client.get("/api/v1/explanation").json()
    assert result["mode"] == "deterministic"
    assert "50.00" in result["text"]
    assert "1,000.00" in result["text"]


PREVIEW_HEADERS = {"content-type": "text/csv", "x-aegisrisk-synthetic": "true"}
CSV_COLUMNS = ["id", "borrower", "industry", "geography", "rating", "pd", "lgd", "ead",
               "currency", "horizon"]
CSV_ROW = ["SYN-CSV-001", "Fictional Orchard", "Energy", "West", "B", "0.1", "0.5",
           "1000.00", "INR", "one_year"]


def csv_body(rows=None, columns=None):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(CSV_COLUMNS if columns is None else columns)
    writer.writerows([CSV_ROW] if rows is None else rows)
    return output.getvalue().encode("utf-8")


def preview(client, body, headers=None):
    return client.post("/api/v1/import-preview", content=body,
                       headers=PREVIEW_HEADERS if headers is None else headers)


def test_preview_is_accurate_and_read_only(client):
    before = client.get("/api/v1/portfolio").json()
    result = preview(client, csv_body()).json()
    assert result["valid"] is True
    assert result["persisted"] is False
    assert result["summary"]["total_expected_loss"] == 50
    assert result["summary"]["total_ead"] == 1000
    assert result["summary"]["borrower_hhi"] == 1
    assert result["row_count"] == result["valid_rows"] == 1
    assert result["issues"] == []
    assert client.get("/api/v1/portfolio").json() == before


def test_preview_utf8_bom_reordered_headers_and_quoted_text(client):
    row = CSV_ROW.copy()
    row[1] = "Fictional Orchard, Ltd"
    result = preview(client, b"\xef\xbb\xbf" + csv_body([row[::-1]], CSV_COLUMNS[::-1]))
    assert result.status_code == 200
    assert result.json()["summary"]["total_expected_loss"] == 50


@pytest.mark.parametrize("field,value", [
    ("pd", "NaN"), ("pd", "Infinity"), ("pd", "1e-3"), ("pd", "1a2"),
    ("pd", "-0.1"), ("pd", "1.01"), ("lgd", "101"), ("pd", "0.12345678901"),
    ("ead", "1000.001"), ("ead", "1000000000001"), ("ead", "1,000"), ("ead", ""),
    ("currency", "USD"), ("horizon", "lifetime"), ("id", "bad/id"),
    ("borrower", "=SUM(A1)"), ("borrower", "@formula"), ("borrower", ""),
    ("borrower", "a" * 101), ("borrower", "two\nlines"), ("rating", "a" * 9),
])
def test_preview_rejects_invalid_rows_without_partial_summary(client, field, value):
    row = CSV_ROW.copy()
    row[0] = "SYN-CSV-002"
    row[CSV_COLUMNS.index(field)] = value
    result = preview(client, csv_body([CSV_ROW, row]))
    assert result.status_code == 422
    data = result.json()
    assert data["valid"] is False
    assert data["summary"] is None
    assert data["invalid_rows"] == 1
    assert any(issue["column"] == field and issue["row"] == 3 for issue in data["issues"])
    assert data["persisted"] is False


def test_preview_duplicate_ids_case_insensitive(client):
    row = CSV_ROW.copy()
    row[0] = " syn-csv-001 "
    data = preview(client, csv_body([CSV_ROW, row])).json()
    assert data["summary"] is None
    assert data["issues"][0]["code"] == "duplicate"


def test_preview_zero_exposure_and_boundary_probabilities(client):
    row = CSV_ROW.copy()
    row[5:8] = ["0", "1", "0"]
    data = preview(client, csv_body([row])).json()
    assert data["valid"] is True
    assert data["summary"]["total_expected_loss"] == 0
    assert data["summary"]["weighted_pd"] is None
    assert data["summary"]["borrower_hhi"] is None


@pytest.mark.parametrize("body", [
    b"", b"\xff", b"id,borrower\nA,B", csv_body([]),
    csv_body(columns=CSV_COLUMNS[:-1] + ["currency"]),
    csv_body() + b'"unterminated', csv_body([CSV_ROW[:-1]]),
    csv_body([CSV_ROW + ["extra"]]),
])
def test_preview_rejects_malformed_files(client, body):
    assert preview(client, body).status_code == 422


def test_preview_requires_acknowledgement_and_content_type(client):
    assert preview(client, csv_body(), {"content-type": "text/csv"}).status_code == 400
    assert preview(client, csv_body(), {"x-aegisrisk-synthetic": "true"}).status_code == 415


def test_preview_size_and_row_limits(client):
    assert preview(client, b"x" * 262_145).status_code == 413
    rows = []
    for index in range(1001):
        row = CSV_ROW.copy()
        row[0] = f"SYN-{index}"
        rows.append(row)
    assert preview(client, csv_body(rows[:1000])).status_code == 200
    assert preview(client, csv_body(rows)).status_code == 413


def test_preview_bounds_error_output(client):
    data = preview(client, csv_body([CSV_ROW] * 102)).json()
    assert data["issue_count"] == 101
    assert len(data["issues"]) == 100
    assert data["issues_truncated"] is True
    assert data["summary"] is None


@pytest.fixture
def saved_client(client, monkeypatch):
    monkeypatch.setenv("SYNTHETIC_IMPORTS_ENABLED", "true")
    return client


def save_csv(client, body=None, **header_overrides):
    body = csv_body() if body is None else body
    headers = PREVIEW_HEADERS | {
        "x-aegisrisk-reviewed": "true",
        "x-aegisrisk-preview-digest": sha256(body).hexdigest(),
    } | header_overrides
    return client.post("/api/v1/saved-portfolios?name=Fictional%20test", content=body,
                       headers=headers)


def test_saving_is_disabled_by_default(client, monkeypatch):
    monkeypatch.delenv("SYNTHETIC_IMPORTS_ENABLED", raising=False)
    assert client.get("/api/v1/saved-portfolios").json()["enabled"] is False
    assert save_csv(client).status_code == 403
    assert preview(client, csv_body()).status_code == 200


def test_save_roundtrip_is_independent_and_idempotent(saved_client):
    client = saved_client
    before = client.get("/api/v1/portfolio").json()
    row = CSV_ROW.copy()
    row[5:8] = ["0.2", "0.5", "3210.99"]
    body = b"\xef\xbb\xbf" + csv_body([row])
    reviewed = preview(client, body).json()
    assert reviewed["digest"] == sha256(body).hexdigest()
    response = save_csv(client, body)
    assert response.status_code == 201
    saved = response.json()["portfolio"]
    assert "csv_text" not in saved
    assert saved["digest"] == reviewed["digest"]
    params = {"dataset_id": saved["id"]}
    selected = client.get("/api/v1/portfolio", params=params).json()
    assert selected["summary"] == reviewed["summary"]
    assert selected["facilities"][0]["id"] == CSV_ROW[0]
    assert selected["facilities"][0]["ead"] == 3210.99
    assert selected["dataset_id"] == saved["id"]
    exposures = client.get("/api/v1/facilities", params=params | {"q": "orchard"}).json()
    assert exposures["total"] == 1
    stress = client.post("/api/v1/stress", params=params,
                         json={"pd_multiplier": 2, "lgd_multiplier": 1}).json()
    assert stress["baseline_el"] == 321.099
    assert stress["stressed_el"] == 642.198
    assert "321.10" in client.get("/api/v1/explanation", params=params).json()["text"]
    retry = save_csv(client, body)
    assert retry.status_code == 200
    assert retry.json()["created"] is False
    assert retry.json()["portfolio"] == saved
    assert client.get("/api/v1/saved-portfolios").json()["total"] == 1
    assert client.get("/api/v1/portfolio").json() == before


@pytest.mark.parametrize("header,value,status", [
    ("x-aegisrisk-reviewed", "false", 400),
    ("x-aegisrisk-synthetic", "false", 400),
    ("x-aegisrisk-preview-digest", "0" * 64, 409),
    ("content-type", "application/json", 415),
])
def test_save_guards_write_nothing(saved_client, header, value, status):
    assert save_csv(saved_client, **{header: value}).status_code == status
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 0


def test_invalid_or_changed_csv_cannot_be_saved(saved_client):
    invalid = CSV_ROW.copy()
    invalid[0] = "SYN-BAD"
    invalid[5] = "2"
    result = save_csv(saved_client, csv_body([CSV_ROW, invalid]))
    assert result.status_code == 422
    assert result.json()["persisted"] is False
    assert save_csv(saved_client, b"x" * 262_145).status_code == 413
    assert save_csv(saved_client, b"not,csv").status_code == 422
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 0


@pytest.mark.parametrize("name", ["", " " * 3, "=formula", "x" * 81, "two\nlines"])
def test_save_rejects_invalid_names(saved_client, name):
    result = saved_client.post("/api/v1/saved-portfolios", params={"name": name},
                               content=csv_body(), headers=PREVIEW_HEADERS)
    assert result.status_code == 422
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 0


def test_saved_list_pagination_and_unknown_ids(saved_client):
    save_csv(saved_client)
    assert saved_client.get("/api/v1/saved-portfolios?offset=1").json()["items"] == []
    assert saved_client.get("/api/v1/saved-portfolios?limit=101").status_code == 422
    for resource in ("portfolio", "facilities", "explanation", "report"):
        assert saved_client.get(f"/api/v1/{resource}?dataset_id={uuid4()}").status_code == 404
        assert saved_client.get(f"/api/v1/{resource}?dataset_id=invalid").status_code == 422
    assert saved_client.post(f"/api/v1/stress?dataset_id={uuid4()}", json={}).status_code == 404


def test_storage_error_leaves_no_partial_snapshot(saved_client, monkeypatch):
    flushed = []
    before = saved_client.get("/api/v1/portfolio").json()

    def broken_commit(session):
        session.flush()
        flushed.append(session.scalar(select(func.count()).select_from(SavedPortfolio)))
        raise SQLAlchemyError("private storage details must not be returned")

    with monkeypatch.context() as patch:
        patch.setattr(Session, "commit", broken_commit)
        response = save_csv(saved_client)
    assert flushed == [1]
    assert response.status_code == 503
    assert response.json() == {"detail": "Data service unavailable"}
    listing = saved_client.get("/api/v1/saved-portfolios").json()
    assert listing["total"] == 0
    assert listing["items"] == []
    assert saved_client.get("/api/v1/portfolio").json() == before
    retry = save_csv(saved_client)
    assert retry.status_code == 201
    assert retry.json()["created"] is True
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 1


def test_corrupt_snapshot_fails_closed(saved_client):
    saved = save_csv(saved_client).json()["portfolio"]
    sessions = app.dependency_overrides[get_session]()
    session = next(sessions)
    try:
        snapshot = session.get(SavedPortfolio, saved["id"])
        snapshot.csv_text = snapshot.csv_text.replace("1000.00", "2000.00")
        session.commit()
    finally:
        sessions.close()
    response = saved_client.get("/api/v1/portfolio", params={"dataset_id": saved["id"]})
    assert response.status_code == 503
    assert response.json() == {"detail": "Saved portfolio integrity check failed"}


def test_report_matches_dashboard_and_is_read_only(client, monkeypatch):
    before = client.get("/api/v1/portfolio").json()
    calls = []
    original = api_module.load_rows

    def tracked_load(session, dataset_id):
        calls.append(dataset_id)
        return original(session, dataset_id)

    monkeypatch.setattr(api_module, "load_rows", tracked_load)
    response = client.get("/api/v1/report")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert calls == [None]
    report = response.json()
    assert report["schema_version"] == "aegisrisk.baseline-report.v1"
    assert report["dataset_id"] is None
    assert report["dataset_version"] == "demo-v1"
    assert report["report_kind"] == "baseline_only"
    assert report["scenario_included"] is False
    assert report["synthetic"] is True
    assert report["currency"] == "INR"
    assert report["horizon"] == "one_year"
    assert datetime.fromisoformat(report["generated_at"]).utcoffset().total_seconds() == 0
    for key in ("summary", "facilities", "segments", "warnings"):
        assert report[key] == before[key]
    inputs = report["source"]["input_rows"]
    assert inputs[0]["pd"] == "0.1"
    assert inputs[0]["lgd"] == "0.5"
    assert inputs[0]["ead"] == "1000"
    canonical = json.dumps(inputs, sort_keys=True, separators=(",", ":"), ensure_ascii=True)
    assert report["source"]["canonical_sha256"] == sha256(canonical.encode()).hexdigest()
    assert Decimal(report["calculation"]["total_expected_loss_decimal"]) == 50
    assert report["calculation"]["decimal_precision"] == 28
    assert "not an LLM" in report["explanation"]["disclaimer"]
    repeated = client.get("/api/v1/report").json()
    del repeated["generated_at"]
    del report["generated_at"]
    assert repeated == report
    assert client.get("/api/v1/portfolio").json() == before


def test_saved_report_canonical_evidence_and_precision(saved_client):
    row = CSV_ROW.copy()
    row[5:8] = ["0.1234567890", "0.9876543210", "3210.99"]
    other = CSV_ROW.copy()
    other[0] = "SYN-CSV-002"
    first = save_csv(saved_client, csv_body([row, other])).json()["portfolio"]
    row[5] = "0.123456789"
    other[7] = "1000.0"
    second = save_csv(saved_client, csv_body([other, row])).json()["portfolio"]
    assert first["id"] != second["id"]
    reports = [saved_client.get("/api/v1/report", params={"dataset_id": item["id"]}).json()
               for item in (first, second)]
    assert reports[0]["source"] == reports[1]["source"]
    for report, item in zip(reports, (first, second), strict=True):
        assert report["dataset_id"] == item["id"]
        assert report["dataset_version"] == "saved-v1:" + item["id"]
        expected = Decimal("0.123456789") * Decimal("0.987654321") * Decimal("3210.99")
        calculation = report["calculation"]
        assert Decimal(calculation["facility_expected_loss_decimal"][0]["expected_loss"]) == expected
        assert Decimal(calculation["total_expected_loss_decimal"]) == expected + 50
        assert "csv_text" not in report["source"]
        selected = saved_client.get("/api/v1/portfolio", params={"dataset_id": item["id"]})
        assert report["summary"] == selected.json()["summary"]


def test_report_respects_disabled_and_corrupt_snapshot_guards(saved_client, monkeypatch):
    saved = save_csv(saved_client).json()["portfolio"]
    params = {"dataset_id": saved["id"]}
    with monkeypatch.context() as patch:
        patch.setenv("SYNTHETIC_IMPORTS_ENABLED", "false")
        assert saved_client.get("/api/v1/report", params=params).status_code == 403
        assert saved_client.get("/api/v1/report").status_code == 200
    sessions = app.dependency_overrides[get_session]()
    session = next(sessions)
    try:
        session.get(SavedPortfolio, saved["id"]).digest = "0" * 64
        session.commit()
    finally:
        sessions.close()
    response = saved_client.get("/api/v1/report", params=params)
    assert response.status_code == 503
    assert "schema_version" not in response.json()


def test_report_handles_empty_and_zero_exposure(client):
    sessions = app.dependency_overrides[get_session]()
    session = next(sessions)
    try:
        session.get(Facility, "A").ead = Decimal(0)
        session.commit()
        zero = client.get("/api/v1/report").json()
        assert zero["summary"]["weighted_pd"] is None
        assert Decimal(zero["calculation"]["total_expected_loss_decimal"]) == 0
        session.execute(delete(Facility))
        session.commit()
        empty = client.get("/api/v1/report").json()
        assert empty["source"]["input_rows"] == []
        assert empty["summary"]["facility_count"] == 0
        assert empty["summary"]["borrower_hhi"] is None
        assert "No facilities loaded." in empty["warnings"]
    finally:
        sessions.close()


@pytest.mark.parametrize("pd,lgd,ead,expected", [
    (1, 1, 1, 50), (2, 1, 1, 100), (10, 10, 1, 1000),
    (0, 1, 1, 0), (1.25, 0.8, 1.5, 75),
])
def test_scenario_report_reconciles_with_stress(client, pd, lgd, ead, expected):
    inputs = {"pd_multiplier": pd, "lgd_multiplier": lgd, "ead_multiplier": ead}
    before = client.get("/api/v1/portfolio").json()
    response = client.post("/api/v1/scenario-report", json=inputs)
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    report = response.json()
    assert report["schema_version"] == "aegisrisk.scenario-report.v1"
    assert report["report_kind"] == "baseline_and_scenario"
    assert report["scenario_included"] is True
    assert report["dataset_id"] is None
    assert report["dataset_version"] == "demo-v1"
    assert report["synthetic"] is True
    assert report["currency"] == "INR"
    assert report["horizon"] == "one_year"
    scenario = report["scenario"]
    live = client.post("/api/v1/stress", json=inputs).json()
    assert {key: scenario[key] for key in live} == live
    assert scenario["inputs"] == inputs | {"industry": None}
    assert scenario["baseline_el"] == 50
    assert scenario["stressed_el"] == expected
    calc = scenario["calculation"]
    assert calc["decimal_precision"] == 28
    assert calc["decimal_rounding"] == "ROUND_HALF_EVEN"
    assert Decimal(calc["baseline_el_decimal"]) == 50
    assert Decimal(calc["stressed_el_decimal"]) == expected
    assert Decimal(calc["delta_el_decimal"]) == expected - 50
    assert calc["affected_facility_count"] == 1
    detail = calc["facilities"][0]
    assert detail["id"] == "A"
    assert Decimal(detail["stressed_el"]) == expected
    assert detail["lgd_capped"] is (lgd == 10)
    baseline = report["baseline"]
    assert baseline["generated_at"] == report["generated_at"]
    assert baseline["scenario_included"] is False
    assert baseline["summary"] == before["summary"]
    assert client.get("/api/v1/portfolio").json() == before


def test_scenario_report_loads_once_and_scopes_saved_rows(saved_client, monkeypatch):
    row = CSV_ROW.copy()
    row[0] = "SYN-OTHER"
    row[2] = "Services"
    row[5:8] = ["0.2", "0.5", "200"]
    saved = save_csv(saved_client, csv_body([row, CSV_ROW])).json()["portfolio"]
    params = {"dataset_id": saved["id"]}
    original = api_module.load_rows
    calls = []

    def tracked_load(session, dataset_id):
        calls.append(str(dataset_id))
        return original(session, dataset_id)

    monkeypatch.setattr(api_module, "load_rows", tracked_load)
    inputs = {"pd_multiplier": 2, "lgd_multiplier": 1, "industry": "Energy"}
    report = saved_client.post(
        "/api/v1/scenario-report", params=params, json=inputs
    ).json()
    assert calls == [saved["id"]]
    assert report["dataset_id"] == report["baseline"]["dataset_id"] == saved["id"]
    assert report["dataset_version"] == "saved-v1:" + saved["id"]
    scenario = report["scenario"]
    assert scenario["baseline_el"] == 70
    assert scenario["stressed_el"] == 120
    assert scenario["delta_el"] == 50
    calc = scenario["calculation"]
    assert calc["affected_facility_count"] == 1
    assert [item["id"] for item in calc["facilities"]] == [CSV_ROW[0], "SYN-OTHER"]
    untouched = calc["facilities"][1]
    assert untouched["affected"] is False
    assert Decimal(untouched["stressed_el"]) == 20
    assert Decimal(untouched["delta_el"]) == 0
    assert sum(Decimal(item["stressed_el"]) for item in calc["facilities"]) == 120
    baseline = saved_client.get("/api/v1/report", params=params).json()
    assert report["baseline"]["source"] == baseline["source"]
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 1


def test_scenario_report_dataset_guards(saved_client, monkeypatch):
    saved = save_csv(saved_client).json()["portfolio"]
    params = {"dataset_id": saved["id"]}
    path = "/api/v1/scenario-report"
    assert saved_client.post(path, params={"dataset_id": str(uuid4())},
                             json={}).status_code == 404
    assert saved_client.post(path, params={"dataset_id": "invalid"},
                             json={}).status_code == 422
    with monkeypatch.context() as patch:
        patch.setenv("SYNTHETIC_IMPORTS_ENABLED", "false")
        assert saved_client.post(path, params=params, json={}).status_code == 403
        assert saved_client.post(path, json={}).status_code == 200
    sessions = app.dependency_overrides[get_session]()
    session = next(sessions)
    try:
        session.get(SavedPortfolio, saved["id"]).digest = "0" * 64
        session.commit()
    finally:
        sessions.close()
    response = saved_client.post(path, params=params, json={})
    assert response.status_code == 503
    assert "schema_version" not in response.json()


def test_scenario_report_zero_baseline_and_empty_portfolio(client):
    sessions = app.dependency_overrides[get_session]()
    session = next(sessions)
    try:
        session.get(Facility, "A").ead = Decimal(0)
        session.commit()
        zero = client.post("/api/v1/scenario-report", json={}).json()["scenario"]
        assert zero["stressed_el"] == 0
        assert zero["delta_pct"] is None
        assert zero["calculation"]["delta_pct_decimal"] is None
        session.execute(delete(Facility))
        session.commit()
        empty = client.post("/api/v1/scenario-report", json={}).json()
        assert empty["scenario"]["calculation"]["facilities"] == []
        assert empty["scenario"]["calculation"]["affected_facility_count"] == 0
        assert empty["baseline"]["source"]["input_rows"] == []
    finally:
        sessions.close()


XLSX_HEADERS = {
    "content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "x-aegisrisk-synthetic": "true",
}


def xlsx_body(rows=None, change=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(CSV_COLUMNS)
    for row in [CSV_ROW] if rows is None else rows:
        sheet.append(row)
    if change is not None:
        change(workbook, sheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_xlsx_preview_numeric_values_and_read_only(client):
    row = CSV_ROW.copy()
    row[5:8] = [0.1, 0.5, 1000]
    raw = xlsx_body([row])
    before = client.get("/api/v1/portfolio").json()
    response = preview(client, raw, XLSX_HEADERS)
    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is True
    assert result["source_format"] == "xlsx"
    assert result["digest"] == sha256(raw).hexdigest()
    assert result["storage_digest"] != result["digest"]
    assert result["summary"]["total_expected_loss"] == 50
    assert result["summary"]["total_ead"] == 1000
    assert result["persisted"] is False
    assert client.get("/api/v1/portfolio").json() == before


@pytest.mark.parametrize("kind", [
    "formula", "boolean", "date", "error", "hidden_row", "hidden_column",
    "extra_sheet", "merged", "hyperlink", "filter", "extra_column", "extra_row",
])
def test_xlsx_rejects_unsupported_features_without_save(saved_client, kind):
    def change(book, sheet):
        if kind in {"formula", "boolean", "date", "error"}:
            sheet["F2"] = {"formula": "=0.1", "boolean": True,
                           "date": datetime(2026, 1, 1), "error": "#DIV/0!"}[kind]
        elif kind == "hidden_row":
            sheet.row_dimensions[2].hidden = True
        elif kind == "hidden_column":
            sheet.column_dimensions["F"].hidden = True
        elif kind == "extra_sheet":
            book.create_sheet("Extra")
        elif kind == "merged":
            sheet.merge_cells("A2:B2")
        elif kind == "hyperlink":
            sheet["B2"].hyperlink = "https://example.invalid/"
        elif kind == "filter":
            sheet.auto_filter.ref = "A1:J2"
        elif kind == "extra_column":
            sheet["K2"] = "unexpected"
        else:
            sheet["A1002"] = "too far"

    raw = xlsx_body(change=change)
    status = 413 if kind in {"extra_column", "extra_row"} else 422
    assert preview(saved_client, raw, XLSX_HEADERS).status_code == status
    headers = XLSX_HEADERS | {"x-aegisrisk-reviewed": "true",
                              "x-aegisrisk-preview-digest": sha256(raw).hexdigest()}
    result = saved_client.post("/api/v1/saved-portfolios?name=Invalid", content=raw,
                               headers=headers)
    assert result.status_code == status
    assert saved_client.get("/api/v1/saved-portfolios").json()["total"] == 0


@pytest.mark.parametrize("part,content,status", [
    ("docProps/custom.xml", b"<!DOCTYPE x [<!ENTITY a 'x'>]><x>&a;</x>", 422),
    ("xl/vbaProject.bin", b"not executable", 422),
    ("../unsafe.xml", b"<x/>", 422),
    ("docProps/oversized.xml", b"<x>" + b"a" * (2 * 1024 * 1024) + b"</x>", 413),
])
def test_xlsx_archive_guards(client, part, content, status):
    output = io.BytesIO(xlsx_body())
    with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr(part, content)
    assert preview(client, output.getvalue(), XLSX_HEADERS).status_code == status


def test_xlsx_corrupt_empty_size_and_row_validation(client):
    assert preview(client, b"not a workbook", XLSX_HEADERS).status_code == 422
    assert preview(client, b"", XLSX_HEADERS).status_code == 413
    assert preview(client, b"x" * 262_145, XLSX_HEADERS).status_code == 413
    assert preview(client, xlsx_body([]), XLSX_HEADERS).status_code == 422
    row = CSV_ROW.copy()
    row[0] = "BAD"
    row[5] = 2
    result = preview(client, xlsx_body([CSV_ROW, row]), XLSX_HEADERS)
    assert result.status_code == 422
    assert result.json()["summary"] is None
    assert result.json()["invalid_rows"] == 1
    assert preview(client, xlsx_body([CSV_ROW, CSV_ROW]), XLSX_HEADERS).status_code == 422


def test_xlsx_save_review_reopen_deduplicate_and_report(saved_client):
    client = saved_client
    raw = xlsx_body()
    reviewed = preview(client, raw, XLSX_HEADERS).json()
    headers = XLSX_HEADERS | {"x-aegisrisk-reviewed": "true",
                              "x-aegisrisk-preview-digest": reviewed["digest"]}
    before = client.get("/api/v1/portfolio").json()
    wrong = client.post("/api/v1/saved-portfolios?name=Wrong", content=raw,
                        headers=headers | {"x-aegisrisk-preview-digest": "0" * 64})
    assert wrong.status_code == 409
    response = client.post("/api/v1/saved-portfolios?name=Excel", content=raw, headers=headers)
    assert response.status_code == 201
    saved = response.json()["portfolio"]
    assert saved["digest"] == reviewed["storage_digest"]
    params = {"dataset_id": saved["id"]}
    assert client.get("/api/v1/portfolio", params=params).json()["summary"] == reviewed["summary"]
    report = client.get("/api/v1/report", params=params).json()
    assert report["dataset_id"] == saved["id"]
    assert report["summary"]["total_expected_loss"] == 50
    scenario = client.post("/api/v1/scenario-report", params=params,
                           json={"pd_multiplier": 2, "lgd_multiplier": 1}).json()
    assert scenario["scenario"]["stressed_el"] == 100
    changed = xlsx_body(change=lambda book, sheet: setattr(book.properties, "title", "Changed"))
    changed_preview = preview(client, changed, XLSX_HEADERS).json()
    assert changed_preview["digest"] != reviewed["digest"]
    assert changed_preview["storage_digest"] == reviewed["storage_digest"]
    assert client.post("/api/v1/saved-portfolios?name=Stale", content=changed,
                       headers=headers).status_code == 409
    retry = client.post("/api/v1/saved-portfolios?name=Retry", content=changed,
                        headers=headers | {"x-aegisrisk-preview-digest":
                                           changed_preview["digest"]})
    assert retry.status_code == 200
    assert retry.json()["created"] is False
    assert retry.json()["portfolio"] == saved
    assert client.get("/api/v1/saved-portfolios").json()["total"] == 1
    assert client.get("/api/v1/portfolio").json() == before


@pytest.mark.parametrize("scenario", [False, True])
def test_pdf_download_uses_same_evidence_and_does_not_write(client, scenario):
    from base64 import b64decode

    from pypdf import PdfReader

    before = client.get("/api/v1/portfolio").json()
    response = (client.post("/api/v1/scenario-report-pdf", json={"pd_multiplier": 2})
                if scenario else client.get("/api/v1/report-pdf"))
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    payload = response.json()
    report = payload["report"]
    content = b64decode(payload["pdf"]["base64"], validate=True)
    assert content.startswith(b"%PDF-")
    assert len(content) == payload["pdf"]["byte_length"]
    assert sha256(content).hexdigest() == payload["pdf"]["sha256"]
    assert payload["pdf"]["media_type"] == "application/pdf"
    pdf = PdfReader(io.BytesIO(content), strict=True)
    text = "\n".join(page.extract_text() for page in pdf.pages)
    baseline = report["baseline"] if scenario else report
    assert baseline["source"]["canonical_sha256"] in text.replace("\n", "")
    assert "Fictional A" in text
    assert "1,000.00" in text and "50.00" in text
    assert "NOT A CREDIT DECISION" in text
    assert len(pdf.pages) >= 2
    for index, page in enumerate(pdf.pages, 1):
        assert f"Page {index}" in page.extract_text()
        assert not page.get("/Annots")
    assert not pdf.trailer["/Root"].get("/OpenAction")
    assert not pdf.trailer["/Root"].get("/AA")
    assert report["scenario_included"] is scenario
    if scenario:
        assert "Scenario comparison" in text
        assert report["scenario"]["inputs"]["pd_multiplier"] == 2
        assert report["scenario"]["stressed_el"] == 110
    else:
        assert "Baseline only." in text
    assert client.get("/api/v1/portfolio").json() == before


@pytest.mark.parametrize("resource", ["report-pdf", "scenario-report-pdf"])
def test_pdf_rejects_missing_or_invalid_dataset(client, resource):
    method = client.post if resource.startswith("scenario") else client.get
    kwargs = {"json": {}} if resource.startswith("scenario") else {}
    assert method(f"/api/v1/{resource}?dataset_id=invalid", **kwargs).status_code == 422
    assert method(f"/api/v1/{resource}?dataset_id={uuid4()}", **kwargs).status_code == 404


def test_pdf_escaping_empty_and_maximum_rows(tmp_path):
    from types import SimpleNamespace

    from pypdf import PdfReader

    from aegisrisk.pdf_report import literal, render_pdf

    assert literal('<img src="https://example.invalid/x"/>') == (
        '&lt;img src="https://example.invalid/x"/&gt;'
    )
    assert literal("Fictional \u03a9") == r"Fictional \u03a9"
    rows = [SimpleNamespace(
        id=f"PDF-{index:04d}", borrower="Fictional " + "x" * 85,
        industry="Energy", geography="West", rating="B",
        pd=Decimal("0.1234567890"), lgd=Decimal("0.9876543210"),
        ead=Decimal("1000000000000.00"),
    ) for index in range(1000)]
    for source in ([], rows):
        report = api_module.baseline_report_payload(source, None)
        content = render_pdf(report)
        pdf = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() for page in pdf.pages)
        if source:
            assert all(row.id in text for row in source)
            assert len(pdf.pages) > 10
            assert len(content) < 4194304
        else:
            assert "No facilities loaded." in text
            assert "Not defined" in text
    with pytest.raises(ValueError, match="facility limit"):
        render_pdf(api_module.baseline_report_payload(rows + [rows[0]], None))


def test_pdf_visual_evidence(tmp_path):
    import os
    from pathlib import Path
    from types import SimpleNamespace

    from pypdf import PdfReader

    from aegisrisk.pdf_report import render_pdf

    target = Path(os.environ.get("AEGISRISK_PDF_EVIDENCE_DIR", str(tmp_path)))
    target.mkdir(parents=True, exist_ok=True)
    rows = [SimpleNamespace(
        id=f"PDF-{i:03d}", borrower=f"Fictional Review Company {i:03d}",
        industry="Energy" if i % 2 else "Technology", geography="West", rating="B",
        pd=Decimal("0.025"), lgd=Decimal("0.45"), ead=Decimal(2500000),
    ) for i in range(40)]
    baseline = api_module.baseline_report_payload(rows, None)
    config = api_module.StressInput(pd_multiplier=2, lgd_multiplier=1.2, industry="Energy")
    scenario = {"report_kind": "baseline_and_scenario", "synthetic": True,
                "generated_at": baseline["generated_at"], "baseline": baseline,
                "scenario": api_module.stress_payload(rows, config, include_evidence=True)}
    for name, report in (("baseline", baseline), ("scenario", scenario)):
        content = render_pdf(report)
        (target / f"{name}.pdf").write_bytes(content)
        reader = PdfReader(io.BytesIO(content))
        assert all(page.extract_text().strip() for page in reader.pages)
    rows[0].borrower = '<img src="https://example.invalid/x"/> & Fictional \u03a9'
    hostile = render_pdf(api_module.baseline_report_payload(rows[:1], None))
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(hostile)).pages)
    assert '<img src="https://example.invalid/x"/>' in text
    assert r"\u03a9" in text
    (target / "literal-labels.pdf").write_bytes(hostile)
